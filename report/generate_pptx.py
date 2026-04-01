"""
Generate a supervisor meeting PPTX presentation.
Cloud-Native AI Operations Agent for CEM-CVM Intelligence
Overview presentation — no code, no ports, just the idea + outputs + diagrams.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Colors ──────────────────────────────────────────────────────────────────
DARK_BG      = RGBColor(0x1A, 0x1A, 0x2E)
HUAWEI_RED   = RGBColor(0xCF, 0x0A, 0x2C)
AI_VIOLET    = RGBColor(0x8E, 0x44, 0xAD)
OSS_BLUE     = RGBColor(0x2E, 0x86, 0xAB)
BSS_GREEN    = RGBColor(0x2E, 0xCC, 0x71)
CEM_TEAL     = RGBColor(0x1A, 0xBC, 0x9C)
CVM_ORANGE   = RGBColor(0xE6, 0x7E, 0x22)
ACTION_RED   = RGBColor(0xE7, 0x4C, 0x3C)
DATALAKE     = RGBColor(0x34, 0x98, 0xDB)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY   = RGBColor(0xBD, 0xC3, 0xC7)
DARK_TEXT     = RGBColor(0x2C, 0x3E, 0x50)
SUBTITLE_CLR = RGBColor(0xAA, 0xAA, 0xBB)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width  = SLIDE_W
prs.slide_height = SLIDE_H

# ── Helpers ─────────────────────────────────────────────────────────────────
def set_dark_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BG

def add_rect(slide, left, top, width, height, fill_color, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape

def set_text(shape, text, size=14, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = alignment
    return tf

def add_paragraph(tf, text, size=14, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, space_before=Pt(4)):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = alignment
    p.space_before = space_before
    return p

def add_textbox(slide, left, top, width, height, text, size=14, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = alignment
    return txBox

def add_bullet_slide_content(slide, bullets, left, top, width, height, size=13, color=WHITE):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_before = Pt(6)
        p.level = 0
    return txBox

def slide_title(slide, title, subtitle=None):
    add_textbox(slide, Inches(0.8), Inches(0.3), Inches(11), Inches(0.7),
                title, size=28, color=WHITE, bold=True)
    if subtitle:
        add_textbox(slide, Inches(0.8), Inches(0.9), Inches(11), Inches(0.5),
                    subtitle, size=14, color=SUBTITLE_CLR)
    # red accent line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(2), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = HUAWEI_RED
    line.line.fill.background()


# =============================================================================
#  SLIDE 1 — TITLE
# =============================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_dark_bg(s)

# ADN L4 badge top-right
badge = add_rect(s, Inches(10.5), Inches(0.4), Inches(2.2), Inches(0.55), AI_VIOLET)
set_text(badge, "ADN Level 4", size=14, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# Title
add_textbox(s, Inches(1.5), Inches(1.8), Inches(10), Inches(1.2),
            "Cloud-Native AI Operations Agent", size=38, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_textbox(s, Inches(1.5), Inches(2.9), Inches(10), Inches(0.8),
            "for CEM\u2013CVM Intelligence", size=32, color=HUAWEI_RED, bold=True, alignment=PP_ALIGN.CENTER)

# Subtitle
add_textbox(s, Inches(2), Inches(4.0), Inches(9), Inches(0.6),
            "An Autonomous Agent That Detects, Decides, and Acts", size=18, color=SUBTITLE_CLR, alignment=PP_ALIGN.CENTER)

# Red line
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(4.8), Inches(4), Pt(3))
line.fill.solid()
line.fill.fore_color.rgb = HUAWEI_RED
line.line.fill.background()

# Author
add_textbox(s, Inches(2), Inches(5.2), Inches(9), Inches(0.5),
            "Souhayl Guenichi  |  ESPRIT \u00d7 Huawei Tunisia  |  March 2026", size=16, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
add_textbox(s, Inches(2), Inches(5.8), Inches(9), Inches(0.5),
            "Supervisor Meeting \u2014 Progress Overview", size=14, color=SUBTITLE_CLR, alignment=PP_ALIGN.CENTER)


# =============================================================================
#  SLIDE 2 — THE PROBLEM
# =============================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_bg(s)
slide_title(s, "The Problem", "Two data worlds that never talk to each other")

# OSS box
oss_box = add_rect(s, Inches(0.8), Inches(1.8), Inches(4.5), Inches(3.2), RGBColor(0x2E, 0x86, 0xAB), OSS_BLUE)
tf = set_text(oss_box, "OSS — Network", size=20, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "", size=8)
add_paragraph(tf, "Throughput", size=14, color=WHITE)
add_paragraph(tf, "Latency", size=14, color=WHITE)
add_paragraph(tf, "Packet Loss", size=14, color=WHITE)
add_paragraph(tf, "Signal Quality (RSRP)", size=14, color=WHITE)
add_paragraph(tf, "Cell Load", size=14, color=WHITE)

# GAP
gap = add_rect(s, Inches(5.6), Inches(2.5), Inches(2), Inches(1.5), HUAWEI_RED)
set_text(gap, "GAP\nNo bridge\nNo action", size=16, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# BSS box
bss_box = add_rect(s, Inches(8.0), Inches(1.8), Inches(4.5), Inches(3.2), RGBColor(0x2E, 0xCC, 0x71), BSS_GREEN)
tf = set_text(bss_box, "BSS — Business", size=20, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "", size=8)
add_paragraph(tf, "Revenue / ARPU", size=14, color=WHITE)
add_paragraph(tf, "Churn Risk", size=14, color=WHITE)
add_paragraph(tf, "Data Usage (DOU)", size=14, color=WHITE)
add_paragraph(tf, "Voice / SMS", size=14, color=WHITE)
add_paragraph(tf, "Subscriber Profiles", size=14, color=WHITE)

# Bottom message
add_textbox(s, Inches(1), Inches(5.5), Inches(11), Inches(0.8),
            "When a cell degrades, the business feels it DAYS later. Nobody connects the dots. Nobody acts.",
            size=16, color=ACTION_RED, bold=True, alignment=PP_ALIGN.CENTER)


# =============================================================================
#  SLIDE 3 — THE SOLUTION
# =============================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_bg(s)
slide_title(s, "Our Solution: The AI Operations Agent", "CEM \u2192 AI Agent \u2192 CVM")

# CEM box
cem = add_rect(s, Inches(0.8), Inches(2.0), Inches(3.0), Inches(2.0), CEM_TEAL)
tf = set_text(cem, "CEM", size=22, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "Huawei SmartCare", size=13, color=WHITE, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "KPI / KQI / CEI", size=12, color=WHITE, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "Experience Data", size=12, color=WHITE, alignment=PP_ALIGN.CENTER)

# Arrow 1
arr1 = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.0), Inches(2.6), Inches(1.0), Inches(0.6))
arr1.fill.solid()
arr1.fill.fore_color.rgb = SUBTITLE_CLR
arr1.line.fill.background()

# Agent box
agent = add_rect(s, Inches(5.2), Inches(1.7), Inches(3.2), Inches(2.8), AI_VIOLET)
tf = set_text(agent, "AI Operations Agent", size=18, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "", size=6)
add_paragraph(tf, "DETECT", size=14, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "Anomalies + SLA Risk", size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "DECIDE", size=14, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "O+B Correlation + Rules", size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "ACT", size=14, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "Recommendations + Triggers", size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# Arrow 2
arr2 = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.6), Inches(2.6), Inches(1.0), Inches(0.6))
arr2.fill.solid()
arr2.fill.fore_color.rgb = SUBTITLE_CLR
arr2.line.fill.background()

# CVM box
cvm = add_rect(s, Inches(9.8), Inches(2.0), Inches(3.0), Inches(2.0), CVM_ORANGE)
tf = set_text(cvm, "CVM", size=22, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "Customer Value Mgmt", size=13, color=WHITE, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "Churn / Upsell / Retain", size=12, color=WHITE, alignment=PP_ALIGN.CENTER)
add_paragraph(tf, "Revenue Impact", size=12, color=WHITE, alignment=PP_ALIGN.CENTER)

# Bottom: differentiators
diffs = ["Real TT Data", "Cloud-Native (HCS)", "O+B Convergence", "Autonomous Actions (L4)"]
for i, d in enumerate(diffs):
    bx = add_rect(s, Inches(1.0 + i * 3.1), Inches(5.3), Inches(2.8), Inches(0.65), RGBColor(0x2C, 0x3E, 0x50), AI_VIOLET)
    set_text(bx, d, size=12, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)


# =============================================================================
#  SLIDE 4 — 5-PHASE PIPELINE
# =============================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_bg(s)
slide_title(s, "How It Works: 5-Phase Pipeline", "From raw operator data to autonomous action")

phases = [
    ("Phase 1", "Data Ingestion", "Ingest real TT data into data lake", CEM_TEAL),
    ("Phase 2", "Feature Engineering", "Build meaningful features from OSS+BSS", OSS_BLUE),
    ("Phase 3", "AI Inference", "SLA risk, anomalies, correlations", AI_VIOLET),
    ("Phase 4", "Decision & Action", "Evaluate rules, generate recommendations", ACTION_RED),
    ("Phase 5", "Curation & Persist", "Store curated data + actions", BSS_GREEN),
]

for i, (phase, name, desc, color) in enumerate(phases):
    y = Inches(1.7) + Inches(i * 1.05)
    box = add_rect(s, Inches(1.5), y, Inches(10), Inches(0.85), color)
    tf = set_text(box, f"{phase}: {name}", size=16, color=WHITE, bold=True)
    add_paragraph(tf, desc, size=12, color=LIGHT_GRAY)
    # arrow between phases
    if i < 4:
        arr = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.3), y + Inches(0.8), Inches(0.4), Inches(0.25))
        arr.fill.solid()
        arr.fill.fore_color.rgb = LIGHT_GRAY
        arr.line.fill.background()


# =============================================================================
#  SLIDE 5 — AI MODELS
# =============================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_bg(s)
slide_title(s, "The Intelligence Layer: AI Models", "Four AI components working together")

models = [
    ("SLA Risk Prediction", "How likely is this region to breach SLA?", "Score 0\u20131 per region", AI_VIOLET),
    ("Network Anomaly Detection", "Is this cell behaving abnormally?", "Anomaly flag + severity", OSS_BLUE),
    ("Revenue Anomaly Detection", "Is subscriber revenue dropping?", "Revenue anomaly flags", BSS_GREEN),
    ("O+B Correlation Engine", "When network degrades, what happens to revenue?", "Pearson + Spearman coefficients", CVM_ORANGE),
]

for i, (title, question, output, color) in enumerate(models):
    col = i % 2
    row = i // 2
    left = Inches(0.8) + Inches(col * 6.2)
    top = Inches(1.7) + Inches(row * 2.4)
    box = add_rect(s, left, top, Inches(5.8), Inches(2.0), color)
    tf = set_text(box, title, size=18, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_paragraph(tf, "", size=6)
    add_paragraph(tf, f'"{question}"', size=13, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
    add_paragraph(tf, f"Output: {output}", size=12, color=WHITE, alignment=PP_ALIGN.CENTER)

# Bottom: arrow to action engine
add_textbox(s, Inches(2), Inches(6.3), Inches(9), Inches(0.5),
            "All outputs feed into the Decision & Action Engine \u2192",
            size=14, color=ACTION_RED, bold=True, alignment=PP_ALIGN.CENTER)


# =============================================================================
#  SLIDE 6 — ACTION ENGINE
# =============================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_bg(s)
slide_title(s, "The Action Engine", "What makes us L4 \u2014 not just analytics, autonomous action")

rules = [
    ("High SLA Risk (> 0.7)", "Proactive Alert", "Critical", ACTION_RED),
    ("OSS Anomaly + Revenue Drop", "Churn Prevention Campaign", "High", HUAWEI_RED),
    ("Healthy Network + Revenue Spike", "Upsell Opportunity", "Medium", BSS_GREEN),
    ("Significant O+B Correlation", "Insight Report for CVM", "Low", DATALAKE),
    ("Multiple Anomalies, Same Region", "Escalation Ticket", "Critical", ACTION_RED),
]

# Header row
hdr = add_rect(s, Inches(0.6), Inches(1.7), Inches(4.5), Inches(0.55), AI_VIOLET)
set_text(hdr, "Trigger Condition", size=13, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
hdr2 = add_rect(s, Inches(5.2), Inches(1.7), Inches(4.5), Inches(0.55), AI_VIOLET)
set_text(hdr2, "Action Generated", size=13, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
hdr3 = add_rect(s, Inches(9.8), Inches(1.7), Inches(2.8), Inches(0.55), AI_VIOLET)
set_text(hdr3, "Priority", size=13, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

for i, (trigger, action, priority, color) in enumerate(rules):
    y = Inches(2.35) + Inches(i * 0.65)
    row_bg = RGBColor(0x22, 0x22, 0x36) if i % 2 == 0 else RGBColor(0x28, 0x28, 0x3E)
    r1 = add_rect(s, Inches(0.6), y, Inches(4.5), Inches(0.55), row_bg)
    set_text(r1, trigger, size=12, color=WHITE, alignment=PP_ALIGN.CENTER)
    r2 = add_rect(s, Inches(5.2), y, Inches(4.5), Inches(0.55), row_bg)
    set_text(r2, action, size=12, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    r3 = add_rect(s, Inches(9.8), y, Inches(2.8), Inches(0.55), color)
    set_text(r3, priority, size=12, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# Lifecycle
add_textbox(s, Inches(0.8), Inches(5.7), Inches(12), Inches(0.5),
            "Action Lifecycle:  Generated  \u2192  Pending  \u2192  Approved / Dismissed  \u2192  Executed",
            size=14, color=SUBTITLE_CLR, alignment=PP_ALIGN.CENTER)
add_textbox(s, Inches(0.8), Inches(6.3), Inches(12), Inches(0.5),
            "The agent decides and acts. Humans handle edge cases only.",
            size=16, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)


# =============================================================================
#  SLIDE 7 — ARCHITECTURE OVERVIEW
# =============================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_bg(s)
slide_title(s, "Architecture Overview", "5 containerised microservices, cloud-portable")

services = [
    ("API Gateway", "REST interface for\nconsuming results", OSS_BLUE),
    ("AI Service", "ML models for\ninference", AI_VIOLET),
    ("Pipeline Worker", "Orchestrates the\n5-phase pipeline", CEM_TEAL),
    ("PostgreSQL", "8-table structured\ndata store", BSS_GREEN),
    ("MinIO", "3-layer data lake\nraw / proc / curated", DATALAKE),
]

for i, (name, desc, color) in enumerate(services):
    left = Inches(0.5) + Inches(i * 2.55)
    box = add_rect(s, left, Inches(1.8), Inches(2.3), Inches(2.2), color)
    tf = set_text(box, name, size=15, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_paragraph(tf, "", size=6)
    add_paragraph(tf, desc, size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# Docker label
docker_box = add_rect(s, Inches(0.3), Inches(4.3), Inches(12.7), Inches(0.55), RGBColor(0x22, 0x22, 0x36), SUBTITLE_CLR)
set_text(docker_box, "Docker Compose  \u2014  runs with one command  \u2014  same images deploy to Huawei Cloud Stack",
         size=13, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# HCS mapping
add_textbox(s, Inches(0.8), Inches(5.2), Inches(5.5), Inches(0.4),
            "Local (Docker)", size=16, color=WHITE, bold=True)
add_textbox(s, Inches(7.0), Inches(5.2), Inches(5.5), Inches(0.4),
            "Huawei Cloud Stack", size=16, color=HUAWEI_RED, bold=True)

mappings = [
    ("Docker containers", "\u2192  ECS / CCE"),
    ("MinIO (object storage)", "\u2192  OBS"),
    ("PostgreSQL", "\u2192  RDS (managed)"),
]
for i, (local, cloud) in enumerate(mappings):
    y = Inches(5.7) + Inches(i * 0.45)
    add_textbox(s, Inches(1.0), y, Inches(5), Inches(0.4), local, size=13, color=LIGHT_GRAY)
    add_textbox(s, Inches(7.2), y, Inches(5), Inches(0.4), cloud, size=13, color=CVM_ORANGE, bold=True)


# =============================================================================
#  SLIDE 8 — DATA
# =============================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_bg(s)
slide_title(s, "Data: Real Operator, Real Results", "Working with real anonymised data from Tunisie Telecom")

bullets_left = [
    "Real production data from Tunisie Telecom",
    "Network data: cell KPIs across regions",
    "Business data: subscriber revenue & usage",
    "Anonymised: hashed IDs, no PII",
    "Dual-mode: real (default) + synthetic (fallback)",
]
bullets_right = [
    "8-table PostgreSQL schema",
    "Full pipeline traceability (run_id)",
    "3-layer data lake in MinIO",
    "Gouvernorat-level location only",
    "Prepaid-dominant market (~80%)",
]

for i, b in enumerate(bullets_left):
    y = Inches(1.8) + Inches(i * 0.7)
    dot = add_rect(s, Inches(0.8), y + Inches(0.08), Inches(0.15), Inches(0.15), CEM_TEAL)
    dot.line.fill.background()
    add_textbox(s, Inches(1.1), y, Inches(5.5), Inches(0.5), b, size=14, color=WHITE)

for i, b in enumerate(bullets_right):
    y = Inches(1.8) + Inches(i * 0.7)
    dot = add_rect(s, Inches(7.0), y + Inches(0.08), Inches(0.15), Inches(0.15), AI_VIOLET)
    dot.line.fill.background()
    add_textbox(s, Inches(7.3), y, Inches(5.5), Inches(0.5), b, size=14, color=WHITE)

# ADN Level label
adn_box = add_rect(s, Inches(3.5), Inches(5.8), Inches(6), Inches(0.7), AI_VIOLET, AI_VIOLET)
set_text(adn_box, "ADN Level 4: AI decides and acts \u2014 humans handle edge cases only",
         size=14, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)


# =============================================================================
#  SLIDE 9 — COMPETITIVE POSITIONING
# =============================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_bg(s)
slide_title(s, "What Makes Us Different", "Competitive positioning against existing solutions")

# Header
cols = ["Solution", "Real Data", "O+B", "Cloud", "Actions"]
widths = [Inches(3.2), Inches(2.0), Inches(2.0), Inches(2.0), Inches(2.5)]
x_positions = [Inches(0.6)]
for w in widths[:-1]:
    x_positions.append(x_positions[-1] + w + Inches(0.12))

for i, (col, w) in enumerate(zip(cols, widths)):
    hdr = add_rect(s, x_positions[i], Inches(1.7), w, Inches(0.5), AI_VIOLET)
    set_text(hdr, col, size=12, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

competitors = [
    ("SELFNET", "\u2717", "\u2717", "\u2713", "\u2717"),
    ("ETSI ZSM", "\u2717", "Partial", "\u2713", "Spec only"),
    ("Nokia AVA", "\u2713", "\u2717", "\u2713", "\u2717"),
    ("SmartCare", "\u2713", "Partial", "\u2713", "\u2717"),
    ("OUR PROJECT", "\u2713", "\u2713", "\u2713", "\u2713"),
]

for r, (name, *vals) in enumerate(competitors):
    y = Inches(2.3) + Inches(r * 0.6)
    is_ours = r == len(competitors) - 1
    bg = AI_VIOLET if is_ours else (RGBColor(0x22, 0x22, 0x36) if r % 2 == 0 else RGBColor(0x28, 0x28, 0x3E))
    font_bold = is_ours

    cell = add_rect(s, x_positions[0], y, widths[0], Inches(0.5), bg)
    set_text(cell, name, size=12, color=WHITE, bold=font_bold, alignment=PP_ALIGN.CENTER)
    for c, val in enumerate(vals):
        cell = add_rect(s, x_positions[c + 1], y, widths[c + 1], Inches(0.5), bg)
        clr = BSS_GREEN if val == "\u2713" else (ACTION_RED if val == "\u2717" else CVM_ORANGE)
        if is_ours:
            clr = WHITE
        set_text(cell, val, size=13, color=clr, bold=font_bold, alignment=PP_ALIGN.CENTER)

add_textbox(s, Inches(1), Inches(5.7), Inches(11), Inches(0.5),
            "We are the only solution combining: Real Data + O+B Convergence + Cloud-Native + Autonomous Actions",
            size=15, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)


# =============================================================================
#  SLIDE 10 — OUTPUTS & DELIVERABLES
# =============================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_bg(s)
slide_title(s, "Outputs & Deliverables So Far", "What we have built and delivered")

deliverables = [
    ("Working Docker Stack", "5 services, runs with one command", CEM_TEAL),
    ("3 Trained ML Models", "GBR + 2x IsolationForest on real TT data", AI_VIOLET),
    ("Decision & Action Engine", "Typed recommendations with priority", ACTION_RED),
    ("3-Layer Data Lake", "Raw, processed, curated in MinIO", DATALAKE),
    ("REST API", "Consume results + manage actions", OSS_BLUE),
    ("8-Table PostgreSQL Schema", "Full pipeline lineage via run_id", BSS_GREEN),
    ("HCS Architecture Design", "Ready for Huawei Cloud deployment", HUAWEI_RED),
    ("Technical Documentation", "Report + architecture docs", CVM_ORANGE),
]

for i, (title, desc, color) in enumerate(deliverables):
    col = i % 2
    row = i // 2
    left = Inches(0.6) + Inches(col * 6.4)
    top = Inches(1.7) + Inches(row * 1.3)
    box = add_rect(s, left, top, Inches(6.0), Inches(1.1), RGBColor(0x22, 0x22, 0x36), color)
    tf = set_text(box, title, size=15, color=color, bold=True)
    add_paragraph(tf, desc, size=12, color=LIGHT_GRAY)


# =============================================================================
#  SLIDE 11 — NEXT STEPS & THANK YOU
# =============================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_bg(s)
slide_title(s, "Next Steps & Questions", "Roadmap and open discussion")

steps = [
    ("Immediate", "Finalise action engine rules with Huawei team feedback", CEM_TEAL),
    ("Short-term", "Deploy to Huawei Cloud Stack (HCS) environment", OSS_BLUE),
    ("Medium-term", "Real-time streaming + advanced ML (LSTM, transformers)", AI_VIOLET),
    ("Long-term", "Reinforcement learning for action optimisation \u2192 ADN L5", HUAWEI_RED),
]

for i, (timeline, step, color) in enumerate(steps):
    y = Inches(1.7) + Inches(i * 0.95)
    tag = add_rect(s, Inches(0.8), y, Inches(2.0), Inches(0.7), color)
    set_text(tag, timeline, size=14, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(s, Inches(3.1), y + Inches(0.12), Inches(9), Inches(0.55),
                step, size=14, color=WHITE)

# Thank you
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(5.5), Inches(9), Pt(2))
line.fill.solid()
line.fill.fore_color.rgb = AI_VIOLET
line.line.fill.background()

add_textbox(s, Inches(1), Inches(5.8), Inches(11), Inches(0.5),
            "Thank you \u2014 Souhayl Guenichi | ESPRIT \u00d7 Huawei Tunisia",
            size=20, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_textbox(s, Inches(1), Inches(6.4), Inches(11), Inches(0.5),
            "An autonomous agent that detects, decides, and acts",
            size=14, color=SUBTITLE_CLR, alignment=PP_ALIGN.CENTER)


# =============================================================================
#  SAVE
# =============================================================================
out_path = "report/supervisor-meeting.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
